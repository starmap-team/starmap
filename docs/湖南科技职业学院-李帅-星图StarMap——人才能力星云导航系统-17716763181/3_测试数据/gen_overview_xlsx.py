# -*- coding: utf-8 -*-
"""生成测试数据总览 Excel（33 组岗位，2026-09-03 公网实测素材）
读: submission/test-data/batch/_manifest.json + 各子目录 output.json
写: submission/test-data/星图StarMap-测试数据总览.xlsx
"""
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BATCH = r"C:\Users\LiShuai\Desktop\Agents\starmap\submission\test-data\batch"
OUT = r"C:\Users\LiShuai\Desktop\Agents\starmap\submission\test-data\星图StarMap-测试数据总览.xlsx"

manifest = json.load(open(os.path.join(BATCH, "_manifest.json"), encoding="utf-8"))

# ── 样式 ──
HEADER_FILL = PatternFill("solid", fgColor="0B1C2C")
HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
TITLE_FONT = Font(name="微软雅黑", size=14, bold=True, color="0B1C2C")
SUB_FONT = Font(name="微软雅黑", size=9, color="687078")
thin = Side(style="thin", color="BECFCC")
BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)
ALT_FILL = PatternFill("solid", fgColor="E8ECEB")


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def style_body(ws, r1, r2, cols):
    for r in range(r1, r2 + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(c >= 4))
            if (r - r1) % 2 == 1:
                cell.fill = ALT_FILL


wb = Workbook()

# ═══════════ Sheet1: 既有岗位能力图谱（27）═══════════
ws = wb.active
ws.title = "既有岗位图谱(27)"
ws.merge_cells("A1:G1")
ws["A1"] = "星图StarMap 测试数据总览 —— 既有岗位能力图谱（27 组）"
ws["A1"].font = TITLE_FONT
ws.merge_cells("A2:G2")
ws["A2"] = "数据源：公网 https://47.120.60.10 GET /api/v1/positions/{岗位名}，2026-09-03 实测 | 每组详见 batch/existing/{#}_{岗位名}/input.json + output.json"
ws["A2"].font = SUB_FONT

headers = ["#", "岗位名称", "领域(domain_tag)", "必备技能数", "加分技能数", "来源总数", "五要素"]
ws.append([])
for i, h in enumerate(headers, 1):
    ws.cell(row=3, column=i, value=h)
style_header(ws, 3, len(headers))

ex = manifest["existing"]
for idx, m in enumerate(ex):
    out = json.load(open(os.path.join(BATCH, "existing", m["dir"], "output.json"), encoding="utf-8"))
    bonus = len(out.get("capability_graph", {}).get("bonus_skills") or [])
    row = [
        idx, m["position"], m["domain_tag"], m["skills"], bonus, m["sources"],
        "✓" if m["five_elements"] else "✗",
    ]
    ws.append(row)
style_body(ws, 4, 3 + len(ex), len(headers))

widths = [5, 38, 16, 10, 10, 10, 8]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A4"

# ═══════════ Sheet2: 新岗位涌现候选（6）═══════════
ws2 = wb.create_sheet("新岗位涌现候选(6)")
ws2.merge_cells("A1:F1")
ws2["A1"] = "星图StarMap 测试数据总览 —— 新岗位（discover 涌现候选，6 组）"
ws2["A1"].font = TITLE_FONT
ws2.merge_cells("A2:F2")
ws2["A2"] = "数据源：POST /api/v1/positions/discover?with_definitions=true（涌现信号+LLM 定义）+ GET /api/v1/positions/{岗位名}（入图详情），2026-09-03 实测"
ws2["A2"].font = SUB_FONT

headers2 = ["#", "岗位名称", "涌现比例", "涌现技能", "discover定义", "入图持久化"]
for i, h in enumerate(headers2, 1):
    ws2.cell(row=3, column=i, value=h)
style_header(ws2, 3, len(headers2))

em = manifest["emerging"]
for idx, m in enumerate(em):
    out = json.load(open(os.path.join(BATCH, "emerging", m["dir"], "output.json"), encoding="utf-8"))
    has_dd = bool(out.get("definition_from_discover", {}).get("industry_scenario"))
    has_pd = bool((out.get("in_graph_detail") or {}).get("persisted_definition"))
    row = [idx, m["position"], m["ratio"], ", ".join(m["emerging_skills"]), "✓" if has_dd else "✗", "✓" if has_pd else "✗"]
    ws2.append(row)
style_body(ws2, 4, 3 + len(em), len(headers2))

widths2 = [5, 44, 10, 34, 12, 12]
for i, w in enumerate(widths2, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A4"

wb.save(OUT)
print(f"已生成: {OUT}")
print(f"Sheet1 既有岗位 {len(ex)} 组 / Sheet2 新岗位 {len(em)} 组")
